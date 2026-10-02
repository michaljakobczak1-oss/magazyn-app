"""
Import katalogu TCL z arkusza „Stan Magazynowy” (TCL Stuff + Archiwum).

Model A:
  - quantity = Przyjęcie TTL (stan bazowy)
  - dostępność liczona rezerwacjami
  - wiersze EVENT / wydane → rezerwacja status=wydane (import)
  - przesunięcia trwałe (inv=0, przes>0, bez powrotu) → archiwum
  - arkusz Archiwum → archived=1
"""
from __future__ import annotations

from datetime import date, datetime, timedelta
from pathlib import Path
import uuid

from openpyxl import load_workbook

from db import (
    get_db, init_db, local_today, replace_equipment_stock, sync_equipment_from_stock,
    sync_equipment_primary_photo,
)


def _cell_str(v):
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%d.%m.%Y")
    if isinstance(v, date):
        return v.strftime("%d.%m.%Y")
    s = str(v).strip()
    if s.lower() in ("n/a", "na", "none", "-"):
        return ""
    return s


def _cell_int(v, default=0):
    if v is None or v == "":
        return default
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return default


def _truthy_flag(v):
    if v is None or v == "":
        return 0
    if isinstance(v, (int, float)):
        return 1 if float(v) != 0 else 0
    s = str(v).strip().lower()
    if s in ("0", "nie", "n", "no", "false", "n/a", "na", "-", "brak", "x"):
        # samotne "x" czasem oznacza "jest" – w tym arkuszu częściej puste = brak
        if s == "x":
            return 1
        return 0
    if s in ("1", "tak", "t", "yes", "true", "jest", "ok"):
        return 1
    return 1


def _norm_header(v):
    if v is None:
        return ""
    return " ".join(str(v).replace("\n", " ").lower().split())


def _find_header_row(ws, max_scan=8):
    for r in range(1, max_scan + 1):
        vals = [_norm_header(ws.cell(r, c).value) for c in range(1, min(30, ws.max_column + 1))]
        joined = " | ".join(vals)
        if "etykieta" in joined and ("model" in joined or "kategoria" in joined):
            return r
    return None


def _map_columns(ws, header_row):
    mapping = {}
    aliases = {
        "etykieta": "code",
        "etykieta quinno": "code",
        "lokalizacja": "location",
        "rząd": "aisle",
        "rzad": "aisle",
        "paleta": "pallet",
        "paleta pal[01-99]": "pallet",
        "kategoria": "category",
        "podkategoria": "subcategory",
        "model": "model",
        "s/n": "serial",
        "sn": "serial",
        "pilot": "remote",
        "kabel zasilający": "cable",
        "kabel": "cable",
        "nogi podstawa": "stand",
        "nogi": "stand",
        "stan (czy działa)": "condition_raw",
        "stan": "condition_raw",
        "przyjęcie ttl": "ttl",
        "wydanie (np. event)": "issued",
        "wydanie": "issued",
        "przesunięcia (nie wraca)": "transferred",
        "przesuniecia (nie wraca)": "transferred",
        "inwentaryzacja (dostępność)": "inventory",
        "inwentaryzacja": "inventory",
        "status:": "status",
        "status": "status",
        "event: kiedy wraca?": "event_return",
        "event": "event_return",
        "kto odebrał": "receiver",
        "kto odebral": "receiver",
        "uwagi:": "notes",
        "uwagi": "notes",
        "ilość": "qty",
        "ilosc": "qty",
        "szczegóły": "details",
        "szczegoly": "details",
        "data": "arch_date",
        "komentarz": "arch_comment",
    }
    for c in range(1, ws.max_column + 1):
        h = _norm_header(ws.cell(header_row, c).value)
        if not h:
            continue
        key = aliases.get(h)
        if not key:
            for a, k in aliases.items():
                if a in h or h in a:
                    key = k
                    break
        if key and key not in mapping:
            mapping[key] = c
    return mapping


def _condition_from_raw(raw, status):
    st = (status or "").strip().upper()
    if st == "UTL":
        return "do utylizacji"
    s = _cell_str(raw).lower()
    if not s:
        return "sprawny"
    if any(x in s for x in ("utyl", "zepsut", "nie działa", "nie dziala", "uszkod")):
        if "utyl" in s:
            return "do utylizacji"
        return "uszkodzony"
    return "sprawny"


def _admin_user_id(con):
    row = con.execute(
        "SELECT id FROM users WHERE role='admin' ORDER BY id LIMIT 1"
    ).fetchone()
    return row["id"] if row else None


def _upsert_equipment(con, *, code, name, location, aisle, pallet, category, subcategory,
                      serial, remote, cable, stand, condition, quantity, notes,
                      archived=False, archived_note=None):
    code = code.strip()
    if not code:
        return None, "empty_code"
    name = (name or code).strip() or code
    existing = con.execute(
        "SELECT id, catalog FROM equipment WHERE code=?", (code,)
    ).fetchone()
    if existing and (existing["catalog"] or "main") != "tcl":
        return None, "other_catalog"

    notes_final = notes or ""
    if archived_note:
        notes_final = (notes_final + "\n" + archived_note).strip() if notes_final else archived_note

    fields = dict(
        project_number=None,
        name=name,
        dimensions=None,
        location=location or None,
        warehouse_id=None,
        owner="TCL",
        brand="TCL",
        material_type="klient",
        condition=condition or "sprawny",
        condition_notes=None,
        storage_instructions=None,
        quantity=max(0, int(quantity or 0)),
        notes=notes_final or None,
        catalog="tcl",
        archived=1 if archived else 0,
        archived_at=local_today().isoformat() if archived else None,
        tcl_category=category or None,
        tcl_subcategory=subcategory or None,
        serial_number=serial or None,
        tcl_aisle=aisle or None,
        tcl_pallet=pallet or None,
        has_remote=int(remote or 0),
        has_power_cable=int(cable or 0),
        has_stand=int(stand or 0),
    )

    if existing:
        eid = existing["id"]
        sets = ", ".join(f"{k}=?" for k in fields)
        con.execute(
            f"UPDATE equipment SET {sets} WHERE id=?",
            (*fields.values(), eid),
        )
        action = "updated"
    else:
        cols = ", ".join(["code", *fields.keys()])
        placeholders = ",".join("?" * (1 + len(fields)))
        cur = con.execute(
            f"INSERT INTO equipment ({cols}) VALUES ({placeholders})",
            (code, *fields.values()),
        )
        eid = cur.lastrowid
        action = "inserted"

    if not archived and fields["quantity"] > 0:
        replace_equipment_stock(con, eid, None, location or "", fields["quantity"])
    else:
        con.execute("DELETE FROM equipment_stock WHERE equipment_id=?", (eid,))
        if archived:
            con.execute(
                "UPDATE equipment SET warehouse_id=NULL, location=? WHERE id=?",
                ((location or "").strip(), eid),
            )
        else:
            sync_equipment_from_stock(con, eid, keep_total=False)
    return eid, action


def _ensure_event_reservation(con, eid, qty, receiver, event_return, notes, admin_id):
    """Tworzy / odświeża rezerwację 'wydane' z importu TCL (model A)."""
    if qty <= 0 or not admin_id:
        return False
    today = local_today()
    # usuń poprzednie importowe rezerwacje TCL dla tej pozycji
    con.execute(
        """DELETE FROM reservations
           WHERE equipment_id=? AND IFNULL(notes,'') LIKE '%[import TCL Stan Magazynowy]%'
             AND status IN ('rezerwacja','wydane')""",
        (eid,),
    )
    date_to = None
    ev = _cell_str(event_return)
    if ev:
        for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y"):
            try:
                date_to = datetime.strptime(ev[:10] if fmt.startswith("%Y") else ev[:10], fmt).date()
                break
            except ValueError:
                try:
                    date_to = datetime.strptime(ev, fmt).date()
                    break
                except ValueError:
                    pass
    # Marek: data wyjścia znana, zwrot często nie – bez daty z Excela: otwarty termin
    date_to_s = date_to.isoformat() if date_to else "9999-12-31"
    note = "[import TCL Stan Magazynowy] Pozycja oznaczona jako EVENT / poza magazynem."
    if notes:
        note += "\n" + notes
    con.execute(
        """INSERT INTO reservations
           (equipment_id, user_id, client, date_from, date_to, quantity, status,
            receiver, recipient_name, notes, issued_at, issued_by)
           VALUES (?,?,?,?,?,?,'wydane',?,?,?,?,?)""",
        (
            eid, admin_id, "TCL / EVENT (import)",
            today.isoformat(), date_to_s, qty,
            (receiver or "")[:120] or None,
            (receiver or "")[:120] or None,
            note,
            today.isoformat(), admin_id,
        ),
    )
    return True


def _row_dict(ws, row, colmap):
    out = {}
    for key, c in colmap.items():
        out[key] = ws.cell(row, c).value
    return out


def _image_ext(data: bytes) -> str:
    if data[:3] == b"\xff\xd8\xff":
        return ".jpeg"
    if data[:8] == b"\x89PNG\r\n\x1a\n":
        return ".png"
    if data[:6] in (b"GIF87a", b"GIF89a"):
        return ".gif"
    if data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return ".webp"
    return ".jpeg"


def _anchor_row_col(img):
    """Zwraca (row_1based, col_1based) albo None."""
    a = img.anchor
    if isinstance(a, str):
        from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
        col_letter, row = coordinate_from_string(a)
        return row, column_index_from_string(col_letter)
    fr = getattr(a, "_from", None)
    if fr is not None:
        return int(fr.row) + 1, int(fr.col) + 1
    return None


def extract_tcl_images_by_code(ws, upload_dir: Path, code_col=1):
    """
    Wyciąga osadzone zdjęcia z arkusza.
    Zwraca {code: {'primary': filename|None, 'extra': [filename, ...]}}.
    Kolumna 7 = zdjęcie główne, 8 = dodatkowe (jak w Stan Magazynowy).
    """
    upload_dir = Path(upload_dir)
    upload_dir.mkdir(parents=True, exist_ok=True)
    by_code = {}
    images = list(getattr(ws, "_images", []) or [])
    for img in images:
        pos = _anchor_row_col(img)
        if not pos:
            continue
        row, col = pos
        code = _cell_str(ws.cell(row, code_col).value)
        if not code:
            continue
        try:
            data = img._data()
        except Exception:
            continue
        if not data:
            continue
        ext = _image_ext(data)
        fname = f"{uuid.uuid4().hex}{ext}"
        (upload_dir / fname).write_bytes(data)
        bucket = by_code.setdefault(code, {"primary": None, "extra": []})
        if col == 7 and not bucket["primary"]:
            bucket["primary"] = fname
        elif col == 8:
            bucket["extra"].append(fname)
        elif col == 7:
            bucket["extra"].append(fname)
        else:
            # inne kolumny (WZ/PZ itd.) – do galerii jako dodatkowe
            bucket["extra"].append(fname)
    return by_code


def attach_photos_to_equipment(con, by_code, replace_existing=True):
    """Podpina wyciągnięte pliki do kart TCL (equipment.photo + equipment_photos)."""
    attached = 0
    for code, files in by_code.items():
        row = con.execute(
            """SELECT id FROM equipment
               WHERE code=? AND IFNULL(catalog,'main')='tcl'""",
            (code,),
        ).fetchone()
        if not row:
            continue
        eid = row["id"]
        primary = files.get("primary")
        extras = list(files.get("extra") or [])
        if not primary and extras:
            primary = extras.pop(0)
        if not primary and not extras:
            continue
        if replace_existing:
            con.execute("DELETE FROM equipment_photos WHERE equipment_id=?", (eid,))
        order = 0
        if primary:
            con.execute(
                """INSERT INTO equipment_photos (equipment_id, filename, sort_order, kind)
                   VALUES (?,?,?,'normal')""",
                (eid, primary, order),
            )
            order += 1
        for fn in extras[:4]:
            con.execute(
                """INSERT INTO equipment_photos (equipment_id, filename, sort_order, kind)
                   VALUES (?,?,?,'normal')""",
                (eid, fn, order),
            )
            order += 1
        sync_equipment_primary_photo(con, eid)
        attached += 1
    return attached


def attach_tcl_photos_from_xlsx(xlsx_path, upload_dir, con=None, sheet_name=None):
    """Samodzielne wyciągnięcie zdjęć z Excela i podpięcie do istniejących kart TCL."""
    path = Path(xlsx_path)
    upload_dir = Path(upload_dir)
    own_con = con is None
    if own_con:
        init_db()
        con = get_db()
    # data_only=False – potrzebne do osadzonych obrazów
    wb = load_workbook(path, data_only=False)
    if not sheet_name:
        for cand in ("TCL Stuff", "TCL", "Sprzęt", "Sheet1"):
            if cand in wb.sheetnames:
                sheet_name = cand
                break
        sheet_name = sheet_name or wb.sheetnames[0]
    ws = wb[sheet_name]
    header_row = _find_header_row(ws) or 3
    colmap = _map_columns(ws, header_row)
    code_col = colmap.get("code", 1)
    by_code = extract_tcl_images_by_code(ws, upload_dir, code_col=code_col)
    n = attach_photos_to_equipment(con, by_code, replace_existing=True)
    con.commit()
    if own_con:
        con.close()
    return {"photos_codes": len(by_code), "attached": n, "sheet": sheet_name}


def run_tcl_stan_import(xlsx_path, con=None, include_archive=True, upload_dir=None):
    """
    Importuje plik Stan Magazynowy.
    Zwraca dict ze statystykami.
    """
    path = Path(xlsx_path)
    if not path.is_file():
        raise FileNotFoundError(str(path))

    own_con = con is None
    if own_con:
        init_db()
        con = get_db()

    if upload_dir is None:
        upload_dir = Path(__file__).resolve().parent / "static" / "uploads"

    wb = load_workbook(path, data_only=True)
    stats = {
        "inserted": 0, "updated": 0, "archived": 0, "events": 0,
        "skipped": 0, "errors": [], "sheet_active": None, "sheet_archive": None,
        "photos_attached": 0, "photos_codes": 0,
    }
    admin_id = _admin_user_id(con)
    active_live_codes = set()  # kody nadal aktywne na TCL Stuff (nie archiwizować hurtowo)

    # --- TCL Stuff ---
    sheet_name = None
    for cand in ("TCL Stuff", "TCL", "Sprzęt", "Sheet1"):
        if cand in wb.sheetnames:
            sheet_name = cand
            break
    if not sheet_name:
        sheet_name = wb.sheetnames[0]
    stats["sheet_active"] = sheet_name
    ws = wb[sheet_name]
    header_row = _find_header_row(ws)
    if not header_row:
        raise ValueError(f"Nie znaleziono wiersza nagłówków w arkuszu {sheet_name}")
    colmap = _map_columns(ws, header_row)
    if "code" not in colmap:
        raise ValueError("Brak kolumny Etykieta w arkuszu aktywnym")

    for r in range(header_row + 1, ws.max_row + 1):
        raw = _row_dict(ws, r, colmap)
        code = _cell_str(raw.get("code"))
        if not code:
            continue
        model = _cell_str(raw.get("model"))
        category = _cell_str(raw.get("category"))
        subcategory = _cell_str(raw.get("subcategory"))
        location = _cell_str(raw.get("location"))
        aisle = _cell_str(raw.get("aisle"))
        pallet = _cell_str(raw.get("pallet"))
        serial = _cell_str(raw.get("serial"))
        status = _cell_str(raw.get("status")).upper()
        ttl = _cell_int(raw.get("ttl"), 0)
        issued = _cell_int(raw.get("issued"), 0)
        transferred = _cell_int(raw.get("transferred"), 0)
        inventory = _cell_int(raw.get("inventory"), -1)
        if inventory < 0:
            inventory = max(0, ttl - issued - transferred)
        if ttl <= 0:
            ttl = max(inventory + issued + transferred, 1 if code else 0)
        notes = _cell_str(raw.get("notes"))
        receiver = _cell_str(raw.get("receiver"))
        event_return = raw.get("event_return")
        condition = _condition_from_raw(raw.get("condition_raw"), status)

        # trwałe przesunięcie całego stanu → archiwum
        permanently_gone = (
            transferred > 0 and inventory <= 0 and issued <= 0 and status != "EVENT"
        ) or status == "UTL" and inventory <= 0 and issued <= 0

        # EVENT / poza magazynem – zostawiamy na stanie, rezerwacja wydane
        out_qty = 0
        if status == "EVENT" or issued > 0:
            out_qty = issued if issued > 0 else max(0, ttl - max(inventory, 0))
            if out_qty <= 0 and inventory <= 0 and ttl > 0:
                out_qty = ttl

        archived = bool(permanently_gone and out_qty <= 0 and status == "UTL")
        # jeśli wszystko przesunięte na stałe bez EVENT – archiwum
        if transferred >= ttl and inventory <= 0 and issued <= 0 and status != "EVENT":
            archived = True

        try:
            eid, action = _upsert_equipment(
                con,
                code=code,
                name=model or subcategory or category or code,
                location=location,
                aisle=aisle,
                pallet=pallet,
                category=category,
                subcategory=subcategory,
                serial=serial,
                remote=_truthy_flag(raw.get("remote")),
                cable=_truthy_flag(raw.get("cable")),
                stand=_truthy_flag(raw.get("stand")),
                condition=condition,
                quantity=0 if archived else ttl,
                notes=notes,
                archived=archived,
            )
            if action == "other_catalog":
                stats["skipped"] += 1
                stats["errors"].append(f"{code}: kod istnieje w katalogu głównym")
                continue
            if action == "empty_code":
                stats["skipped"] += 1
                continue
            if action == "inserted":
                stats["inserted"] += 1
            else:
                stats["updated"] += 1
            if archived:
                stats["archived"] += 1
            else:
                active_live_codes.add(code)
                if out_qty > 0 and eid:
                    if _ensure_event_reservation(
                        con, eid, min(out_qty, ttl), receiver, event_return, notes, admin_id
                    ):
                        stats["events"] += 1
        except Exception as e:
            stats["skipped"] += 1
            stats["errors"].append(f"{code}: {e}")

    # --- Archiwum ---
    # Uwaga: w Excelu ten sam kod bywa i na TCL Stuff (np. TTL 20 / inv 8),
    # i w Archiwum (notatki o sztukach, które odeszły). Nie wolno wtedy
    # nadpisać całej karty jako archived – zostaje aktywna wg TCL Stuff.
    if include_archive and "Archiwum" in wb.sheetnames:
        stats["sheet_archive"] = "Archiwum"
        stats["archive_notes_only"] = 0
        wa = wb["Archiwum"]
        hr = _find_header_row(wa)
        if hr:
            amap = _map_columns(wa, hr)
            for r in range(hr + 1, wa.max_row + 1):
                raw = _row_dict(wa, r, amap)
                code = _cell_str(raw.get("code"))
                if not code:
                    continue
                model = _cell_str(raw.get("model"))
                qty = _cell_int(raw.get("qty"), 1)
                details = _cell_str(raw.get("details"))
                comment = _cell_str(raw.get("arch_comment"))
                arch_date = _cell_str(raw.get("arch_date"))
                note_parts = [p for p in (
                    details, comment,
                    f"Data archiwum: {arch_date}" if arch_date else "",
                    "[import Archiwum Stan Magazynowy]",
                ) if p]
                note_blob = "\n".join(note_parts)

                existing = con.execute(
                    """SELECT id, quantity, IFNULL(archived,0) AS archived, notes
                       FROM equipment
                       WHERE code=? AND IFNULL(catalog,'main')='tcl'""",
                    (code,),
                ).fetchone()

                # Kod nadal żywy na TCL Stuff → nie archiwizuj i NIE doklejaj
                # tekstów z arkusza Archiwum do aktywnej karty.
                if code in active_live_codes:
                    stats["archive_notes_only"] = stats.get("archive_notes_only", 0) + 1
                    continue

                try:
                    eid, action = _upsert_equipment(
                        con,
                        code=code,
                        name=model or _cell_str(raw.get("subcategory")) or code,
                        location="",
                        aisle="",
                        pallet="",
                        category=_cell_str(raw.get("category")),
                        subcategory=_cell_str(raw.get("subcategory")),
                        serial=_cell_str(raw.get("serial")),
                        remote=0, cable=0, stand=0,
                        condition="sprawny",
                        quantity=qty,
                        notes="\n".join([p for p in (details, comment, f"Data archiwum: {arch_date}" if arch_date else "") if p]),
                        archived=True,
                        archived_note="[import Archiwum Stan Magazynowy]",
                    )
                    if action in ("inserted", "updated"):
                        if action == "inserted":
                            stats["inserted"] += 1
                        else:
                            stats["updated"] += 1
                        stats["archived"] += 1
                    elif action == "other_catalog":
                        stats["skipped"] += 1
                except Exception as e:
                    stats["skipped"] += 1
                    stats["errors"].append(f"Archiwum {code}: {e}")

    # Zdjęcia osadzone w Excelu (wymaga ponownego otwarcia bez data_only)
    try:
        photo_stats = attach_tcl_photos_from_xlsx(
            path, upload_dir, con=con, sheet_name=stats["sheet_active"]
        )
        stats["photos_codes"] = photo_stats.get("photos_codes", 0)
        stats["photos_attached"] = photo_stats.get("attached", 0)
    except Exception as e:
        stats["errors"].append(f"Zdjęcia: {e}")

    con.commit()
    if own_con:
        con.close()
    return stats


if __name__ == "__main__":
    import json
    import sys
    src = sys.argv[1] if len(sys.argv) > 1 else str(
        Path.home() / "Downloads" / "Stan Magazynowy-4.xlsx"
    )
    print("Import:", src)
    result = run_tcl_stan_import(src)
    print(json.dumps({k: v for k, v in result.items() if k != "errors"}, ensure_ascii=False, indent=2))
    if result["errors"]:
        print("errors (max 20):")
        for e in result["errors"][:20]:
            print(" -", e)
